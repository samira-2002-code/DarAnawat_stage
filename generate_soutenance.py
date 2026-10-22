from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from docx import Document
from docx.shared import Inches as DI, Pt as DP, RGBColor as DR
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\HP\Desktop\Nouveau dossier (2)\DarAnawat_stage")
OUT = ROOT / "soutenance_reservation"
OUT.mkdir(exist_ok=True)
SESSION = Path(r"C:\Users\HP\.copilot\session-state\07965710-3f24-4071-bf53-a74363c53c24\files")
FORM_RAW = SESSION / "reservation-form-real.png"
OK_RAW = SESSION / "reservation-success-real.png"
FORM = OUT / "capture-formulaire-reelle.png"
OK = OUT / "capture-confirmation-reelle.png"

def crop(src, dst, bottom):
    im = Image.open(src).convert("RGB")
    im.crop((0, 0, im.width, min(bottom, im.height))).save(dst)

crop(FORM_RAW, FORM, 2500)
crop(OK_RAW, OK, 2600)

# Editable PowerPoint: text and diagrams are native PowerPoint objects.
P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
BLANK = P.slide_layouts[6]
COL = {"forest":"12332F","deep":"0D2522","gold":"C99D58","palegold":"E4C78E",
       "cream":"F7F4ED","white":"FFFEFA","ink":"18312D","muted":"68756E",
       "line":"E4DCCB","mint":"E7F0E8","red":"9B5148","pale":"EFE9DE"}
def color(v): return RGBColor.from_string(COL.get(v,v))
def bg(s, c="cream"):
    s.background.fill.solid(); s.background.fill.fore_color.rgb=color(c)
def box(s,x,y,w,h,c,rounded=False,line=None,trans=0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                          Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=color(c); sh.fill.transparency=trans
    if line: sh.line.color.rgb=color(line); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    return sh
def txt(s,v,x,y,w,h,size=14,c="ink",bold=False,font="Aptos",align=None,val=MSO_ANCHOR.TOP,margin=.03):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(margin)
    tf.vertical_anchor=val
    for i,line in enumerate(v.split("\n")):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.space_after=Pt(1)
        if align is not None: p.alignment=align
        for r in p.runs:
            r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color(c)
    return sh
def pill(s,v,x,y,w,c="gold",fg="deep"):
    box(s,x,y,w,.34,c,True)
    txt(s,v.upper(),x+.08,y+.075,w-.16,.17,8.3,fg,True,align=PP_ALIGN.CENTER)
def footer(s,n,src=""):
    txt(s,"DAR ANAWAT  /  MODULE DEMANDE DE RÉSERVATION",.62,7.13,6.2,.18,8,"muted",True)
    if src: txt(s,src,6.55,7.13,5.55,.18,7.2,"muted",align=PP_ALIGN.RIGHT)
    txt(s,f"{n:02}",12.62,7.1,.25,.2,9,"gold",True,align=PP_ALIGN.RIGHT)
def head(s,k,t,n,src=""):
    txt(s,k.upper(),.68,.4,8,.23,9,"gold",True)
    txt(s,t,.66,.77,12,.72,27,"ink",True,"Georgia")
    footer(s,n,src)
def card(s,x,y,w,h,title,body,dark=False,accent="gold"):
    box(s,x,y,w,h,"forest" if dark else "white",True)
    box(s,x,y,.06,h,accent,True)
    txt(s,title,x+.23,y+.18,w-.42,.38,14.5,"white" if dark else "ink",True)
    txt(s,body,x+.23,y+.66,w-.42,h-.78,10.8,"palegold" if dark else "muted")
def code(s,x,y,w,h,label,body,fs=9.6):
    box(s,x,y,w,h,"deep",True); pill(s,label,x+.2,y+.17,min(1.6,w-.4))
    txt(s,body,x+.22,y+.66,w-.44,h-.79,fs,"palegold",font="Consolas")

# 1. Context and scope
s=P.slides.add_slide(BLANK); bg(s,"deep")
box(s,.68,.62,.52,.52,"gold",True); txt(s,"DA",.73,.79,.42,.15,11,"deep",True,align=PP_ALIGN.CENTER)
pill(s,"Soutenance · contribution personnelle",1.42,.72,3.55)
txt(s,"Dar Anawat",.72,1.7,5.9,.56,22,"palegold",True,"Georgia")
txt(s,"Du séjour rêvé\nà la demande enregistrée",.68,2.36,7.3,1.52,31,"white",True,"Georgia")
txt(s,"Mon périmètre : relier le formulaire de demande\nau traitement Laravel et à la persistance.",.74,4.2,6,.8,16,"palegold")
txt(s,"Site de riad · découverte des chambres · demande de séjour à vérifier par l’équipe.",.74,5.5,5.9,.63,12,"white")
box(s,8.15,1.55,4.35,4.68,"forest",True)
txt(s,"LE FLUX, EN UNE IMAGE",8.55,1.98,3.3,.25,10,"gold",True)
for i,(n,label) in enumerate([("01","FORMULAIRE"),("02","VALIDATION"),("03","DEMANDE PENDING")]):
    yy=2.58+i*1.03; box(s,8.55,yy,.48,.48,"gold",True)
    txt(s,n,8.66,yy+.15,.25,.12,8,"deep",True,align=PP_ALIGN.CENTER)
    txt(s,label,9.24,yy+.11,2.85,.25,13,"white",True)
    if i<2: box(s,8.77,yy+.56,.035,.44,"gold")
txt(s,"Pas de confirmation automatique. Pas de paiement en ligne.",8.55,5.85,3.4,.34,10.5,"palegold")
txt(s,"GIT : commits du module de réservation attribués à Samira Assemlali.",.74,6.65,9.9,.22,9,"palegold")
footer(s,1,"README.md · routes/web.php")

# 2. Need
s=P.slides.add_slide(BLANK); bg(s); head(s,"01 · Le besoin","Une demande claire, sans promesse de disponibilité",2,"README.md · lang/fr/site.php")
txt(s,"Le module transforme une intention de séjour en demande structurée, exploitable par l’équipe.",.72,1.75,11.1,.45,16,"muted")
card(s,.75,2.55,3.65,2.55,"Le visiteur renseigne","Ses coordonnées\nLa chambre et la capacité\nLes dates et ses demandes particulières")
card(s,4.82,2.55,3.65,2.55,"Le système garantit","Les champs obligatoires\nDes dates cohérentes\nLe respect de la capacité choisie")
card(s,8.9,2.55,3.65,2.55,"L’équipe reprend la main","La demande est enregistrée « pending »\nDisponibilité à vérifier\nAucun paiement en ligne",accent="red")
box(s,.75,5.55,11.8,.78,"forest",True)
txt(s,"POSITIONNEMENT EXACT",1.02,5.8,2.3,.2,9,"gold",True)
txt(s,"Une demande de séjour — pas un moteur d’inventaire ni une réservation confirmée.",3.3,5.72,8.65,.35,14,"white",True)

# 3. UX, with real browser capture
s=P.slides.add_slide(BLANK); bg(s); head(s,"02 · L’expérience","Un parcours guidé, dans l’identité du site",3,"reservations/create.blade.php")
txt(s,"COORDONNÉES",.82,1.82,2.1,.22,9,"gold",True)
txt(s,"Nom · e-mail · téléphone facultatif",.82,2.17,4.2,.32,14,"ink",True)
txt(s,"SÉJOUR",.82,2.82,2.1,.22,9,"gold",True)
txt(s,"Type de chambre · voyageurs\nArrivée · départ",.82,3.17,4,.7,15,"ink",True)
txt(s,"OPTIONNEL",.82,4.1,2.1,.22,9,"gold",True)
txt(s,"Demandes particulières",.82,4.44,3.9,.32,14,"ink",True)
txt(s,"Sélection préremplie depuis le détail d’une chambre ; capacité et dates aussi ajustées côté navigateur.",.82,5.18,4.4,.82,12,"muted")
box(s,5.55,1.62,4,4.93,"pale",True)
txt(s,"INTERACTION",5.92,2.05,2.8,.23,9,"gold",True)
txt(s,"1 chambre\n→ capacité maximale\n\n1 date d’arrivée\n→ départ au moins le lendemain",5.92,2.48,3.1,1.73,14,"ink",True)
txt(s,"Côté client = aide à la saisie.\nLa validation faisant autorité est côté Laravel.",5.92,4.8,3.15,.9,11.5,"muted")
im=Image.open(FORM); ratio=im.width/im.height; h=5.85; w=h*ratio
box(s,10.22,1.18,w+.26,6.08,"line",True)
s.shapes.add_picture(str(FORM),Inches(10.35),Inches(1.3),height=Inches(h))
txt(s,"CAPTURE RÉELLE · FORMULAIRE",10.28,6.72,2.5,.18,8,"muted",True)

# 4. HTTP/data flow
s=P.slides.add_slide(BLANK); bg(s); head(s,"03 · Le parcours des données","Du navigateur au message de confirmation",4,"routes/web.php · ReservationController.php")
steps=[("GET","/reservation","create()\nrenvoie la vue"),("SAISIE","Blade","formulaire +\njeton CSRF"),("POST","/reservation","route web\nthrottle 10/min"),("GARDE","FormRequest","valide et\nrejette sinon"),("ÉCRITURE","Eloquent","Reservation::create\nstatus=pending"),("RETOUR","Redirect","flash succès +\nréférence")]
for i,(tag,title,body) in enumerate(steps):
    x=.68+i*2.12
    box(s,x,2.45,1.78,2.32,"white",True); pill(s,tag,x+.17,2.68,.88 if len(tag)<6 else 1.15)
    txt(s,title,x+.17,3.21,1.46,.44,13.2,"ink",True,"Consolas" if title.startswith("/") else "Aptos")
    txt(s,body,x+.17,3.8,1.46,.72,10.3,"muted")
    if i<5:
        sh=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x+1.83),Inches(3.4),Inches(.25),Inches(.32))
        sh.fill.solid(); sh.fill.fore_color.rgb=color("gold"); sh.line.fill.background()
box(s,.85,5.45,11.65,.8,"forest",True)
txt(s,"ÉCHANGE HTTP CLASSIQUE",1.15,5.73,2.75,.22,9,"gold",True)
txt(s,"Pas d’AJAX : POST → redirection → message de session.",4,5.67,8,.34,13,"white",True)

# 5. Backend
s=P.slides.add_slide(BLANK); bg(s); head(s,"04 · Le backend","Laravel valide, puis persiste ce qui est autorisé",5,"routes · StoreReservationRequest · contrôleur")
code(s,.72,1.8,5.72,3.6,"ROUTE + CONTRÔLEUR","Route::post('/reservation',\n [ReservationController::class, 'store'])\n ->middleware('throttle:10,1')\n ->name('reservations.store');\n\nReservation::create([\n ...$request->validated(),\n 'user_id' => $request->user()?->getAuthIdentifier(),\n 'status' => 'pending',\n]);",9.1)
code(s,6.78,1.8,5.8,3.6,"VALIDATION SERVEUR","'guest_email' => ['required','email','max:254'],\n'room_type' => ['required', Rule::in(\n array_keys(Reservation::ROOM_CAPACITIES))],\n'check_in' => ['required','date_format:Y-m-d',\n 'after_or_equal:today'],\n'check_out' => ['required','date_format:Y-m-d',\n 'after:check_in'],\n// after(): guests <= capacité",8.9)
card(s,.78,5.75,3.63,.82,"Route nommée","Le formulaire POST vers la route.",True)
card(s,4.84,5.75,3.63,.82,"FormRequest","Règles avant store().",True)
card(s,8.9,5.75,3.63,.82,"Contrôleur","pending serveur, redirect + référence.",True)

# 6. Schema/relations
s=P.slides.add_slide(BLANK); bg(s); head(s,"05 · La base de données","Une relation utilisateur facultative, une demande conservée",6,"Reservation.php · User.php · migration")
box(s,.9,2.05,3.45,3.55,"forest",True); txt(s,"users",1.2,2.34,2.7,.38,21,"white",True,"Consolas")
txt(s,"PK  id\nname · email · password\n\nUn utilisateur peut avoir\nplusieurs demandes.",1.2,2.96,2.85,1.55,12,"palegold")
box(s,5.15,2.05,6.95,3.55,"white",True,"line"); txt(s,"reservations",5.48,2.34,5.8,.38,21,"ink",True,"Consolas")
txt(s,"PK id   FK user_id NULL → users.id\nguest_name · guest_email · guest_phone?\nroom_type · guests · check_in · check_out\nspecial_requests? · status · timestamps",5.48,2.96,5.9,1.48,12,"muted",font="Consolas")
ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(4.35),Inches(3.82),Inches(5.15),Inches(3.82)); ln.line.color.rgb=color("gold"); ln.line.width=Pt(2)
txt(s,"1",4.25,3.47,.2,.22,11,"gold",True,align=PP_ALIGN.CENTER); txt(s,"0..*",4.62,3.47,.5,.22,11,"gold",True,align=PP_ALIGN.CENTER)
txt(s,"Index : e-mail + (room_type, status, check_in, check_out) · FK user_id ON DELETE SET NULL",1,5.95,11,.4,11,"muted",True)
txt(s,"room_type est un code texte, pas une FK vers une table rooms.",1,6.43,10.8,.27,10,"red",True)

# 7. Validation/security/test status
s=P.slides.add_slide(BLANK); bg(s); head(s,"06 · Fiabilité","Les règles métier sont vérifiées côté serveur",7,"FormRequest · Blade · tests/Feature/ReservationTest.php")
card(s,.75,1.85,5.72,1.65,"DATES & CAPACITÉ","Arrivée >= aujourd’hui ; départ > arrivée ; voyageurs <= capacité.")
card(s,6.83,1.85,5.72,1.65,"ERREURS EXPLICITES","Retour au formulaire avec erreurs et anciennes valeurs ; Blade affiche erreurs/succès.")
card(s,.75,3.88,5.72,1.65,"PROTECTIONS PRÉSENTES","@csrf · throttle:10,1 · champs validés · statut serveur · échappement Blade.")
card(s,6.83,3.88,5.72,1.65,"TESTS FONCTIONNELS","7 cas / 27 assertions : stockage, capacité, dates, compte, entrée malformée, chevauchements.")
box(s,.78,5.94,11.7,.6,"pale",True)
txt(s,"Résultat local : 0 échec · 27 assertions · 7 avertissements Dotenv (.env absent, clé temporaire).",1,6.14,11.1,.22,10,"muted",True)

# 8. Demo proof
s=P.slides.add_slide(BLANK); bg(s); head(s,"07 · Démonstration vérifiée","La référence confirme l’enregistrement — pas la disponibilité",8,"Capture réelle + ligne SQLite de démonstration")
txt(s,"CAS DE DÉMO",.82,1.92,2.5,.23,9,"gold",True)
txt(s,"Chambre Double\n2 voyageurs · 1–3 nov. 2026\nDemande test « Demo Guest »",.82,2.34,4.7,1.2,18,"ink",True)
box(s,.82,3.92,5.3,.75,"mint",True); txt(s,"RÉFÉRENCE 1  ·  STATUS pending",1.12,4.17,4.8,.24,13,"forest",True)
txt(s,"SELECT id, room_type, check_in, check_out,\n       guests, status\nFROM reservations WHERE id = 1;",.82,5,6.15,.95,10.5,"ink",font="Consolas")
txt(s,"SQLite temporaire isolée pour cette préparation ; pas une base de production.",.82,6.3,6.2,.45,9.5,"muted")
im=Image.open(OK); h=5.8
s.shapes.add_picture(str(OK),Inches(9.4),Inches(1.35),height=Inches(h))
txt(s,"CONFIRMATION AFFICHÉE PAR L’APPLICATION",8.45,6.72,4.2,.18,8,"muted",True,align=PP_ALIGN.CENTER)

# 9. Contribution, limits and perspectives
s=P.slides.add_slide(BLANK); bg(s,"deep")
txt(s,"08 · BILAN",.72,.5,4,.22,9,"gold",True)
txt(s,"Périmètre livré.\nSuite logique : gérer le stock.",.68,.95,8,1.15,28,"white",True,"Georgia")
card(s,.75,2.55,3.65,2.2,"MA CONTRIBUTION","Migration + modèle\nRoutes + FormRequest\nContrôleur\nVue formulaire + tests",True)
card(s,4.84,2.55,3.65,2.2,"LIMITE ACTUELLE","Pas d’inventaire\nPas de confirmation UI\nPas de paiement\nChevauchements acceptés",True,"red")
card(s,8.92,2.55,3.65,2.2,"PISTES, NON LIVRÉES","Disponibilité / inventaire\nWorkflow équipe\nNotifications\nRègles de concurrence",True)
txt(s,"Git : Samira Assemlali — commits 47bc6d0, 0ea514e, 6060bb5, 302b82f. Les commits antérieurs du frontend général sont attribués à saraFront.",.82,5.5,11.65,.76,10.5,"palegold")
txt(s,"À retenir : le module recueille une demande structurée ; l’équipe vérifie la disponibilité.",.82,6.48,11.5,.32,13,"white",True)
footer(s,9,"README.md · historique Git")

PPTX=OUT/"DarAnawat_Soutenance_Reservation.pptx"; P.save(PPTX)

# Word document helpers
def setup_doc(d):
    sec=d.sections[0]; sec.top_margin=DI(.65); sec.bottom_margin=DI(.65); sec.left_margin=DI(.78); sec.right_margin=DI(.78)
    d.styles["Normal"].font.name="Aptos"; d.styles["Normal"].font.size=DP(10)
    d.styles["Normal"].font.color.rgb=DR.from_string("18312D")
    for n,z,c in [("Title",26,"12332F"),("Heading 1",18,"12332F"),("Heading 2",13,"C99D58"),("Heading 3",11,"18312D")]:
        st=d.styles[n]; st.font.name="Aptos Display"; st.font.size=DP(z); st.font.bold=True; st.font.color.rgb=DR.from_string(c)
def label(d,k,v):
    p=d.add_paragraph(); p.paragraph_format.space_after=DP(4)
    r=p.add_run(k+" "); r.bold=True; r.font.color.rgb=DR.from_string("C99D58"); p.add_run(v)
def bullet(d,v):
    p=d.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=DP(2); p.add_run(v)

slides=[
("Dar Anawat — du séjour rêvé à la demande enregistrée",
"Dar Anawat\nDu séjour rêvé à la demande enregistrée\nMa responsabilité : le module de demande de réservation\nSite de riad · découverte des chambres · demande à vérifier",
"Couverture vert profond/or, flux éditable formulaire → validation → pending.",
"Dar Anawat est un site de maison d’hôtes qui permet de découvrir le lieu et ses chambres. Ma présentation se limite à ma partie : le parcours de demande de séjour, du formulaire jusqu’à l’enregistrement. Le mot important est demande : le système ne promet pas la disponibilité et ne prend pas de paiement.",
"Une demande de contact structurée n’est pas une réservation confirmée.",
"README.md décrit la maison d’hôtes et les chambres. lang/fr/site.php précise l’attente de vérification et l’absence de paiement. Le Git attribue les commits ciblés réservation à Samira Assemlali.",
"Que fait Dar Anawat ? — Un site de présentation de riad et de chambres, avec une demande de séjour ; la confirmation reste hors du flux automatique."),
("Une demande claire, sans promesse de disponibilité",
"Coordonnées + séjour + demandes particulières\nValidation serveur des données\nDemande enregistrée « pending »\nDisponibilité vérifiée ensuite par l’équipe",
"Trois cartes : saisie, garanties du système, rôle de l’équipe.",
"Le module réunit les informations utiles pour répondre : qui voyage, quel type de chambre, à quelles dates et avec quelle capacité. La validation garantit la cohérence des données. Elle ne sait pas si une chambre est déjà occupée, car il n’y a pas d’inventaire connecté.",
"Validation = contrôle de complétude/cohérence, pas recherche de disponibilité.",
"lang/fr/site.php affiche explicitement que la disponibilité doit être vérifiée et qu’aucun paiement en ligne n’a lieu.",
"Pourquoi « demande » et non « réservation » ? — Le statut créé est pending ; un test autorise deux demandes chevauchantes."),
("Un parcours guidé, dans l’identité du site",
"Vos coordonnées : nom, e-mail, téléphone facultatif\nDétails du séjour : chambre, voyageurs, arrivée, départ\nDemandes particulières facultatives\nAide de saisie et erreurs affichées au champ",
"Capture réelle du formulaire sur mobile, repères sur les groupes et les aides.",
"Le formulaire sépare les coordonnées des détails du séjour. Le téléphone et le commentaire sont facultatifs. Si on arrive depuis le détail d’une chambre, son type est présélectionné. Le JavaScript ajuste la capacité et le départ minimal, mais le serveur refait ces contrôles.",
"Blade produit le HTML côté serveur ; JavaScript améliore l’usage, pas la sécurité.",
"create.blade.php contient guest_name, guest_email, guest_phone, room_type, guests, check_in, check_out, special_requests, @csrf et les attributs data-room-select/data-check-in.",
"Pourquoi revalider au serveur ? — Le navigateur peut être modifié ou contourné."),
("Du navigateur au message de confirmation",
"GET /reservation → affichage du formulaire\nPOST /reservation → validation Laravel\nÉcriture Reservation::create(...)\nRedirect → message flash + référence",
"Chaîne HTTP en six blocs ; mention de soumission classique, pas AJAX.",
"À l’ouverture, la route GET appelle create et renvoie la vue Blade. À l’envoi, le navigateur fait un POST vers la route nommée. Le middleware limite les envois, le FormRequest valide et le contrôleur enregistre. Le contrôleur redirige vers le formulaire avec la référence.",
"Une route associe méthode HTTP et URL à un traitement : GET affiche, POST transmet une action.",
"routes/web.php définit GET /reservation → create et POST /reservation → store, avec throttle:10,1. La vue utilise route('reservations.store').",
"À quoi sert le nom de route ? — À éviter les URL codées en dur et à générer les liens dans vues et tests."),
("Laravel valide, puis persiste ce qui est autorisé",
"POST nommé + throttle:10,1\nFormRequest : règles et capacité\nContrôleur : user éventuel\nStatut fixé : pending",
"Deux blocs de code : route/contrôleur et validation ; badges de responsabilités.",
"Le contrôleur ne recode pas toutes les règles : StoreReservationRequest intervient avant store. rules valide les champs et after compare le nombre de voyageurs à la capacité. Le contrôleur ne prend que validated(), associe l’utilisateur connecté s’il existe, fixe pending et sauvegarde avec Eloquent.",
"FormRequest valide ; contrôleur orchestre ; modèle Eloquent mappe l’objet à la table.",
"StoreReservationRequest vérifie e-mail, dates, Rule::in et capacité. ReservationController::store utilise ...$request->validated() et impose status=pending.",
"Comment empêcher status=confirmed ? — Le contrôleur n’utilise pas le champ reçu et fixe lui-même pending."),
("Une relation utilisateur facultative, une demande conservée",
"users 1 — 0..* reservations\nuser_id nullable → users.id\nroom_type = code texte\nIndex e-mail + index composé séjour/état",
"ERD éditable users → reservations, clés, champs et index.",
"Une demande peut être anonyme, donc user_id est nullable. Un utilisateur peut avoir plusieurs demandes. Si le compte est supprimé, nullOnDelete conserve la demande et retire le lien. room_type est une valeur texte, pas une FK vers rooms.",
"Migration = schéma versionné ; modèle = représentation PHP et relations.",
"La migration définit FK nullable, ON DELETE SET NULL, dates, e-mail indexé et index composite. Reservation::user() est BelongsTo ; User::reservations() est HasMany.",
"Pourquoi nullOnDelete ? — Conserver la demande quand un compte disparaît."),
("Les règles métier sont vérifiées côté serveur",
"Dates ordonnées · capacité contrôlée\nErreurs + anciennes valeurs rendues au formulaire\nCSRF · throttle · statut serveur · échappement Blade\n7 tests / 27 assertions ; avertissement .env absent",
"Quatre cartes validation, retour d’erreur, protections, couverture ; note d’environnement.",
"Les dates passées et durées nulles sont refusées ; le nombre de voyageurs respecte la chambre. En erreur, Laravel redirige avec erreurs et valeurs anciennes ; en succès, la session affiche une référence. Le POST est limité, le formulaire a CSRF et Blade échappe les sorties. Sept tests : 27 assertions sans échec ; l’absence de .env produit des avertissements Dotenv.",
"CSRF protège les formulaires ; throttle ralentit les répétitions ; échappement Blade réduit XSS.",
"@csrf dans la vue ; throttle:10,1 sur le POST ; tests/Feature/ReservationTest.php couvre stockage, capacité, dates, compte, entrée malformée et chevauchements.",
"Quels tests ? — Formulaire, lien depuis la chambre, sauvegarde pending, compte, dates, capacité, type malformé et chevauchement."),
("La référence confirme l’enregistrement — pas la disponibilité",
"Chambre Double · 2 voyageurs\n1–3 novembre 2026 · référence 1\nLigne SQLite : status = pending\nDemo Guest est fictif",
"Capture réelle de succès et mini-sortie SQL issue de la base temporaire.",
"Je remplis avec des données de démonstration, je soumets puis montre la référence. Je vérifie ensuite la ligne reservations : chambre double, deux personnes, statut pending. La SQLite de cette préparation est temporaire ; ce n’est pas une base de production.",
"La référence est la clé primaire ; le message flash survit au redirect dans la session.",
"Le navigateur a affiché « Votre demande (référence 1) a bien été enregistrée ». La ligne vérifiée contient Demo Guest, double, 2026-11-01 → 2026-11-03, 2, pending.",
"Est-ce une chambre réservée ? — Non, seulement une demande enregistrée ; la disponibilité est à vérifier."),
("Périmètre livré. Suite logique : gérer le stock.",
"Livré : formulaire → validation → pending\nLimites : pas d’inventaire, de décision UI, ni de paiement\nPistes : disponibilité, workflow équipe, notifications\nAttribution Git : Samira Assemlali",
"Trois colonnes : contribution, limites actuelles, pistes futures non livrées.",
"Ma contribution couvre persistance, modèle, routes, validation, contrôleur, formulaire et tests. L’interface générale des chambres existait dans des commits antérieurs et je ne la revendique pas. La suite logique est de modéliser le stock et de vérifier les conflits avant confirmation.",
"Séparer l’existant des perspectives évite de présenter des maquettes comme des fonctions livrées.",
"Commits 47bc6d0, 0ea514e, 6060bb5, 302b82f ; le test overlapping_requests_remain_pending garde les chevauchements pending.",
"Que développer ensuite ? — Inventaire, contrôle transactionnel de disponibilité, puis outil d’acceptation/refus ; ce serait une évolution.")
]

D=Document(); setup_doc(D)
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run("DAR ANAWAT"); r.bold=True; r.font.size=DP(13); r.font.color.rgb=DR.from_string("C99D58")
p=D.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Préparer ma soutenance\nModule de demande de réservation")
p=D.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Guide oral · explications techniques · démonstration · questions du jury")
D.add_heading("Résumé vérifiable du module",1)
D.add_paragraph("Dar Anawat est présenté dans le README comme un site de maison d’hôtes : accueil, découverte des chambres et demandes de séjour. Ma partie analysée est uniquement le formulaire et son traitement Laravel. Il collecte coordonnées, chambre, dates, voyageurs et demandes facultatives ; une demande valide est enregistrée au statut pending avec une référence. Ce n’est ni une confirmation de chambre ni un paiement.")
for b in [
"Frontend : vue Blade resources/views/reservations/create.blade.php, style crème/vert/or, responsive, erreurs/succès et aides JavaScript.",
"Flux : GET reservations.create → Blade ; POST reservations.store → StoreReservationRequest → ReservationController@store → Eloquent → redirect avec message flash.",
"Règles : identité/e-mail requis ; chambre single/double/suite ; arrivée aujourd’hui ou plus tard ; départ après arrivée ; voyageurs selon capacité 1/2/4.",
"Données : reservations ; user_id nullable avec FK ON DELETE SET NULL ; index e-mail et index composite chambre/statut/dates ; User hasMany / Reservation belongsTo.",
"Protections constatées : CSRF, validation serveur, statut fixé côté serveur, échappement Blade, throttle:10,1 sur POST. authorize() retourne true : le formulaire est public.",
"Non livré : inventaire/calendrier, blocage des conflits, confirmation/refus en back-office, paiement ou notification métier dans ce flux."
]: bullet(D,b)
D.add_heading("Contribution personnelle et attribution",1)
D.add_paragraph("Le log local attribue à Samira Assemlali les commits 47bc6d0 (schéma/persistance), 0ea514e (routes, FormRequest, contrôleur et traductions de réservation), 6060bb5 (formulaire responsive) et 302b82f (tests et documentation). Les commits antérieurs du frontend général sont attribués à saraFront/sara devlop7. Formulation prudente : « Dans l’historique de cette branche, les commits du module réservation sont attribués à mon identité Git, Samira Assemlali ; le frontend antérieur est attribué à saraFront. » Vérifier que cette identité correspond bien à vous avant l’oral.")
D.add_heading("Tests constatés",1)
D.add_paragraph("Exécution : php artisan test --filter=ReservationTest avec APP_KEY temporaire (aucun .env créé). Résultat : 7 tests, 27 assertions, aucune assertion en échec. PHPUnit émet néanmoins 7 avertissements Dotenv car .env est absent. Ne pas annoncer « zéro avertissement ».")
D.add_page_break(); D.add_heading("Plan oral détaillé — une fiche par diapositive",1)
for i,(title,display,visual,script,concept,proof,q) in enumerate(slides,1):
    if i>1: D.add_page_break()
    D.add_heading(f"Diapositive {i} — {title}",1)
    label(D,"Texte exact à afficher :",display); label(D,"Visuel :",visual)
    label(D,"Script oral naturel :",script); label(D,"Concept simplement :",concept)
    label(D,"Exemple réel dans le dépôt :",proof); label(D,"Question possible :",q)
    if i==3:
        D.add_picture(str(FORM),width=DI(1.55)); D.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    if i==8:
        D.add_picture(str(OK),width=DI(1.55)); D.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER

D.add_page_break(); D.add_heading("Démonstration réaliste — 2 à 3 minutes",1)
D.add_paragraph("Le scénario a été rejoué sur une base SQLite temporaire isolée, avec des données fictives. Pour le jour J, vérifier la configuration DB réellement utilisée par l’application.")
for k,v in [
("1 · Présenter","Ouvrir /reservation ; montrer coordonnées, chambre, voyageurs, arrivée/départ, demande facultative. Le détail d’une chambre peut préselectionner le type."),
("2 · Remplir","Données rejouées : Demo Guest, demo@example.test, Chambre Double, 2 voyageurs, 1–3 novembre 2026. Données fictives."),
("3 · Soumettre","Cliquer Envoyer la demande : POST HTML classique vers reservations.store, protégé par CSRF. Ne pas dire « réservation confirmée »."),
("4 · Montrer le backend","Ouvrir routes/web.php, StoreReservationRequest et ReservationController. Aucun écran admin ne fait partie de ce parcours."),
("5 · Vérifier la base","SELECT id,user_id,guest_name,room_type,check_in,check_out,guests,status FROM reservations WHERE id=1;. Sortie vérifiée : id 1, user_id NULL, double, dates, 2, pending."),
("6 · Résultat","Montrer la référence 1 et préciser que l’équipe doit encore vérifier la disponibilité.")
]: label(D,k,v)
D.add_heading("Si l’environnement n’est pas prêt",2)
D.add_paragraph("Ne fabriquez pas de résultat. Montrer la capture réelle fournie et préciser qu’elle provient d’une soumission locale sur SQLite temporaire. Avant une démo live : vérifier APP_KEY, migrations, connexion à la base et utiliser des données fictives.")

QA=[
("Quel est l’objectif de Dar Anawat ?","Présenter une maison d’hôtes et ses chambres ; le site permet une demande de séjour."),
("Qui utilise le formulaire ?","Un visiteur potentiel, connecté ou non ; le formulaire est public."),
("Demande ou réservation ?","Demande enregistrée pending ; disponibilité non contrôlée et séjour non confirmé."),
("Rôle d’une route Laravel ?","Associer verbe HTTP et URL à une action ; GET affiche, POST transmet."),
("Pourquoi nommer les routes ?","Utiliser route(...) et les tests sans répéter d’URL en dur."),
("Rôle du contrôleur ?","Orchestrer les données validées, l’utilisateur éventuel, la création et la redirection."),
("Pourquoi un FormRequest ?","Centraliser autorisation et règles avant la méthode store."),
("Pourquoi valider côté serveur ?","Les champs HTML/JavaScript sont contournables ; le serveur fait autorité."),
("Comment vérifier la capacité ?","ROOM_CAPACITIES : single=1, double=2, suite=4 ; after() compare au nombre demandé."),
("Peut-on envoyer confirmed ?","Le contrôleur impose pending et ne prend que validated()."),
("Rôle du modèle Reservation ?","Mappe la table, permet assignation contrôlée, cast dates et relation belongsTo."),
("Rôle de la migration ?","Décrire de façon versionnée les colonnes, types, FK et index."),
("Relation SQL ?","FK nullable reservations.user_id → users.id ; User hasMany Reservation."),
("Pourquoi nullOnDelete ?","Garder la demande si le compte associé est supprimé."),
("Les conflits sont-ils refusés ?","Non ; le test confirme que des chevauchements restent pending."),
("Où sont contrôlées les dates ?","Dans le FormRequest ; la migration n’ajoute pas de CHECK SQL de ces règles."),
("Protections présentes ?","CSRF, throttle, validation, statut serveur et échappement Blade ; pas de chiffrement par champ identifié."),
("Que voit-on en cas d’erreur ?","Redirect avec erreurs et anciennes valeurs en session, affichées par Blade."),
("Pourquoi redirect après POST ?","Éviter une page POST au rafraîchissement et afficher le flash sur GET."),
("Quels tests exécutés ?","7 tests/27 assertions sans échec ; avertissements Dotenv dus au .env absent."),
("Pourquoi le téléphone est facultatif ?","Le formulaire et les règles le déclarent nullable ; l’e-mail et le nom sont requis."),
("Quelles langues ?","FR, EN, ES, AR ; layout RTL en arabe."),
("Amélioration prioritaire ?","Inventaire et contrôle transactionnel des disponibilités, puis workflow d’équipe."),
("Quid des données personnelles ?","Nom et e-mail sont stockés ; pas de politique de rétention ni chiffrement applicatif par champ identifié dans le module.")
]
D.add_page_break(); D.add_heading("Questions probables du jury — réponses simples",1)
for q,a in QA: D.add_heading(q,2); D.add_paragraph(a)
D.add_page_break(); D.add_heading("Mémo technique : expliquer sans réciter",1)
for k,v in [
("Route","Point d’entrée HTTP : GET create affiche la page ; POST store reçoit et lance l’action."),
("FormRequest","authorize décide si la requête est permise ; rules valide les champs ; after applique la règle croisée capacité/voyageurs."),
("Contrôleur","Coordonne la réponse HTTP, l’écriture Eloquent et le redirect ; il n’est pas le validateur."),
("Modèle Eloquent","Objet PHP représentant une demande ; fillable autorise l’assignation de masse ; relations vers User."),
("Migration","Historique reproductible du schéma SQL ; colonnes, nullable, FK et index."),
("Relation","user_id référence users.id ; nullable autorise une demande anonyme."),
("Session","Les erreurs, old values et le message flash survivent à la redirection."),
("Sécurité","Validation, CSRF, throttle, statut serveur, échappement Blade ; ces protections ne couvrent pas toutes les exigences de confidentialité.")
]: label(D,k,v)
D.add_heading("Fichiers source à ouvrir pendant la préparation",1)
for f in ["routes/web.php","app/Http/Controllers/ReservationController.php","app/Http/Requests/StoreReservationRequest.php","app/Models/Reservation.php","app/Models/User.php","database/migrations/2026_10_19_000000_create_reservations_table.php","resources/views/reservations/create.blade.php","tests/Feature/ReservationTest.php","README.md","lang/fr/site.php"]: bullet(D,f)
GUIDE=OUT/"Guide_Oral_DarAnawat_Reservation.docx"; D.save(GUIDE)

# Compact printable revision sheet
Q=Document(); setup_doc(Q)
sec=Q.sections[0]; sec.top_margin=DI(.45); sec.bottom_margin=DI(.45); sec.left_margin=DI(.58); sec.right_margin=DI(.58)
Q.styles["Normal"].font.size=DP(9)
Q.add_paragraph("FICHE EXPRESS · RÉSERVATION",style="Title")
Q.add_paragraph("Dar Anawat | Laravel · Blade · Eloquent")
Q.add_heading("La phrase à garder",1)
p=Q.add_paragraph(); p.add_run("Le formulaire enregistre une demande de séjour validée au statut pending ; il ne réserve pas une chambre et ne prend aucun paiement.").bold=True
Q.add_heading("Flux — 6 mots",1)
Q.add_paragraph("GET /reservation → Blade → POST reservations.store → FormRequest → Reservation::create(status=pending) → redirect + référence.")
Q.add_heading("Champs et règles",1)
for b in ["Requis : nom, e-mail, chambre, arrivée, départ, voyageurs. Facultatifs : téléphone, demandes particulières.",
"Capacités : single=1 · double=2 · suite=4. Arrivée >= aujourd’hui ; départ > arrivée ; voyageurs <= capacité.",
"Le JavaScript aide la saisie ; le serveur revalide. user_id est facultatif et lié si connecté."]: bullet(Q,b)
Q.add_heading("Laravel en une ligne",1)
for k,v in [("Route","URL + verbe → action"),("FormRequest","règles avant contrôleur"),("Contrôleur","orchestre, impose pending, écrit, redirige"),("Modèle","Eloquent + relations"),("Migration","schéma SQL versionné"),("FK","users 1 → 0..* reservations ; user_id nullable")]:
    label(Q,k+":",v)
Q.add_heading("Sécurité · erreurs · limites",1)
for b in ["@csrf · throttle:10,1 · données validated() · statut imposé · sortie Blade échappée.",
"Erreur : retour avec messages/anciennes valeurs. Succès : redirect + flash + référence.",
"Chevauchements non bloqués ; disponibilité vérifiée par l’équipe. Pas d’inventaire, back-office ou paiement dans ce flux."]: bullet(Q,b)
Q.add_heading("Tests et attribution",1)
Q.add_paragraph("7 tests / 27 assertions, aucune assertion en échec ; 7 avertissements Dotenv car .env absent. Git attribue 47bc6d0, 0ea514e, 6060bb5, 302b82f à Samira Assemlali. Vérifier que l’identité correspond à vous.")
Q.add_heading("5 questions probables",1)
for q,a in [("Pourquoi pending ?","Disponibilité non vérifiée."),("Pourquoi FormRequest ?","Règles centralisées avant contrôleur."),("Pourquoi user_id nullable ?","Demande publique sans compte."),("Conflit refusé ?","Non, chevauchements restent pending."),("Suite logique ?","Inventaire + contrôle atomique + workflow équipe.")]:
    p=Q.add_paragraph(); p.paragraph_format.space_after=DP(1); r=p.add_run(q+" "); r.bold=True; p.add_run(a)
Q.add_heading("Avant la démo",1)
Q.add_paragraph("Vérifier APP_KEY, DB, migrations et connexion. Utiliser des données fictives. Lire la ligne dans la base configurée. Si elle n’est pas prête, présenter la capture et dire qu’elle vient d’une SQLite temporaire.")
QUICK=OUT/"Fiche_Revision_Rapide_Reservation.docx"; Q.save(QUICK)

print("Fichiers créés :")
for p in [PPTX,GUIDE,QUICK,FORM,OK]: print(p.name, p.stat().st_size, "octets")
print("Diapositives :",len(P.slides), "| Questions jury :",len(QA))
